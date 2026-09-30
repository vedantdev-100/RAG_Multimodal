"""
DoclingParser against real Docling conversions (markdown and DOCX need no
model downloads). The regression these guard: tables and pictures were
silently DROPPED by the first parser because Docling's TableItem and
PictureItem have no `.text` attribute.
"""
import io

import pytest
from docling_core.types.doc.document import PictureDescriptionData, PictureItem

from app.exceptions import IngestionError, ModelNotFoundError
from app.rag.ingestion.parsers.docling_parser import DoclingParser, flatten_docling_document

pytestmark = pytest.mark.asyncio

MARKDOWN = b"""# Sales Report

Revenue grew in Q3.

## Results

| Quarter | Revenue | Growth |
|---------|---------|--------|
| Q1      | 100     | 5%     |
| Q2      | 120     | 20%    |

![Revenue chart](chart.png)

Closing paragraph after the figure.
"""


@pytest.fixture(scope="module")
def parser():
    return DoclingParser()


async def test_table_is_kept_with_its_values_and_heading_path(parser):
    parsed = await parser.parse(MARKDOWN, "report.md")

    tables = [e for e in parsed.elements if e.modality == "table"]
    assert len(tables) == 1
    assert "Q2" in tables[0].text and "120" in tables[0].text and "20%" in tables[0].text
    assert tables[0].headings == ["Sales Report", "Results"]
    assert parsed.metadata["tables"] == 1
    assert parsed.native is not None  # HybridChunker needs the rich document


async def test_picture_is_counted_and_reported_as_not_searchable_without_description(parser):
    parsed = await parser.parse(MARKDOWN, "report.md")
    assert parsed.metadata["pictures"] == 1
    assert parsed.metadata["pictures_described"] == 0


async def test_picture_description_becomes_searchable_text(parser):
    parsed = await parser.parse(MARKDOWN, "report.md")
    picture = next(i for i, _ in parsed.native.iterate_items() if isinstance(i, PictureItem))
    picture.annotations.append(
        PictureDescriptionData(text="Bar chart: revenue rose from 100 in Q1 to 120 in Q2.", provenance="test")
    )

    elements, stats = flatten_docling_document(parsed.native)

    images = [e for e in elements if e.modality == "image"]
    assert len(images) == 1
    assert "revenue rose from 100" in images[0].text
    assert images[0].headings == ["Sales Report", "Results"]
    assert stats["pictures_described"] == 1


async def test_docx_table_and_picture_are_extracted(parser):
    from docx import Document
    from PIL import Image

    document = Document()
    document.add_heading("Quarterly Summary", level=1)
    document.add_paragraph("Numbers below.")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Quarter", "Revenue"
    table.rows[1].cells[0].text, table.rows[1].cells[1].text = "Q2", "120"
    png = io.BytesIO()
    Image.new("RGB", (60, 60), "red").save(png, format="PNG")
    png.seek(0)
    document.add_picture(png)
    buffer = io.BytesIO()
    document.save(buffer)

    parsed = await parser.parse(buffer.getvalue(), "summary.docx")

    assert parsed.metadata["tables"] == 1
    assert parsed.metadata["pictures"] == 1
    table = next(e for e in parsed.elements if e.modality == "table")
    assert "Revenue" in table.text and "120" in table.text


async def test_plain_text_is_parsed(parser):
    parsed = await parser.parse(b"Just some plain text.\n\nA second paragraph.", "notes.txt")
    assert any("second paragraph" in e.text for e in parsed.elements)


async def test_corrupt_file_becomes_a_clean_ingestion_error(parser):
    with pytest.raises(IngestionError) as exc:
        await parser.parse(b"this is not a zip file", "broken.docx")
    assert "broken.docx" in str(exc.value)
    assert "/tmp" not in str(exc.value) and "\\Temp" not in str(exc.value)  # no server paths leak


async def test_pdf_without_downloaded_models_gives_actionable_error(parser):
    with pytest.raises(ModelNotFoundError) as exc:
        await parser.parse(b"%PDF-1.4", "scan.pdf")
    assert "app.cli.download_models" in str(exc.value)
